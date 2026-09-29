"""H.1 (PLAN_VOLTAJES_EDITABLES_BRAQUI_28-09.md, tercera ronda): el
formulario y el PDF dicen el formato correcto.

CAUSA (medida): en `data/widgets.xlsx`, las hojas `encabezado_mensualBraqui`,
`encabezado_ActividadBraqui` y `encabezado_LinealidadBraqui` tenían
`code_1 = "IDC-F-RT-126"` y `code_2 = "V.02"`: el código del formato
**diario**. Los formatos vigentes (dejados por el físico el 28-09) son
**IDC-F-RT-143 V.03** (actividad: mensual, cambio de fuente y redundante) e
**IDC-F-RT-144 V.03** (linealidad). `codigo_de_formato` lleva ese texto al
encabezado de la pantalla y a `id_maquina` del PDF.

El código y la versión son DATOS de la interfaz (una celda del Excel), no
código: el cambio vive en su único sitio. Este archivo prueba (a) lo que
dicen las pantallas construidas, (b) lo que recibe el PDF y (c) que el
diario y los demás formatos siguen diciendo lo que decían -- por camino
del usuario, sobre una copia real de `Rebuild_21-09`."""
import os
import shutil

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt5.QtWidgets import QApplication, QMessageBox

import data.ManejoDatos.conection as conection_mod
import ui.paginasControles.PruebasDiarias.braquiterapia as braq_mod
import ui.paginasControles.PruebasMensuales.braq_mensual as braq_mensual_mod
from ui.paginasControles.PruebasDiarias.braquiterapia import (
    CalRedundanteFuente, Linealidad, PruebaDiariaBraq,
)
from ui.paginasControles.PruebasMensuales.braq_mensual import PruebaMensualBraq
from ui.util_formato import codigo_de_formato
from _bd_referencia import BD_REBUILD_21_09

WIDGETS_XLSX = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "widgets.xlsx")

CODIGO_ACTIVIDAD = "IDC-F-RT-143 V.03"
CODIGO_LINEALIDAD = "IDC-F-RT-144 V.03"
CODIGO_DIARIO = "IDC-F-RT-126 V.02"

# La decisión declarada: a qué formato oficial pertenece cada hoja de
# encabezado. Sale de los Excel del servicio; cambiarla es una decisión,
# no un descuido -- por eso el censo es exacto (ni una hoja de más ni una
# de menos).
CODIGOS_POR_HOJA = {
    "encabezado_600": ("IDC-F-RT-118", "V.04"),
    "encabezado_ix": ("IDC-F-RT-118", "V.04"),
    "encabezado_braq": ("IDC-F-RT-126", "V.02"),
    "encabezado_mensualBraqui": ("IDC-F-RT-143", "V.03"),
    "encabezado_ActividadBraqui": ("IDC-F-RT-143", "V.03"),
    "encabezado_LinealidadBraqui": ("IDC-F-RT-144", "V.03"),
    "encabezado_mensu_600": ("IDC-F-RT-119", "V.03"),
    "encabezado_mensu_IX": ("IDC-F-RT-119", "V.03"),
    "encabezado_mensu_Halcyon": ("IDC-F-RT-119", "V.03"),
    "encabezado_anual_600": ("IDC-F-RT-119", "V.03"),
    "encabezado_anual_ix": ("IDC-F-RT-119", "V.03"),
    "encabezado_anual_Halcyon": ("IDC-F-RT-119", "V.03"),
}


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def _sin_dialogos_modales(monkeypatch):
    for tipo in ("information", "warning", "critical", "question"):
        monkeypatch.setattr(QMessageBox, tipo, staticmethod(lambda *a, **k: None))


@pytest.fixture
def bd_copia_real(monkeypatch, tmp_path):
    """Copia REAL de `Rebuild_21-09` -- nunca se conecta directo a la BD de
    referencia (regla de sesión)."""
    if not BD_REBUILD_21_09.exists():
        pytest.skip(f"falta la BD de referencia {BD_REBUILD_21_09.name}")
    copia = tmp_path / "copia.db"
    shutil.copy2(BD_REBUILD_21_09, copia)
    monkeypatch.setattr(conection_mod, "ruta_base_datos", lambda: str(copia))
    conection_mod.Conexion._instance = None
    yield copia
    if conection_mod.Conexion._instance is not None:
        conection_mod.Conexion._instance.con.close()
    conection_mod.Conexion._instance = None


class _UsuarioFalso:
    def __init__(self, nombre="Físico de Prueba"):
        self._nombre = nombre
        self.nombre = nombre
        self.id = 1

    def __getattr__(self, item):
        return None


def _codigos_del_excel(ruta=WIDGETS_XLSX):
    """{hoja: (code_1, code_2)} de cada hoja que declara los dos widgets:
    la PRIMERA columna `nombres` (la que `createInterface` carga) y la
    columna `descripcion`, que es el texto que se muestra."""
    resultado = {}
    wb = openpyxl.load_workbook(ruta, data_only=True)
    for ws in wb.worksheets:
        encabezado = [c.value for c in ws[1]]
        if "nombres" not in encabezado or "descripcion" not in encabezado:
            continue
        i_nombre = encabezado.index("nombres")
        i_desc = encabezado.index("descripcion")
        valores = {}
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if i_nombre < len(fila) and fila[i_nombre] in ("code_1", "code_2"):
                valores[fila[i_nombre]] = fila[i_desc]
        if "code_1" in valores or "code_2" in valores:
            resultado[ws.title] = (valores.get("code_1"), valores.get("code_2"))
    return resultado


# ---------------------------------------------------------------------------
# (a) Lo que dicen las pantallas construidas.
# ---------------------------------------------------------------------------

class TestElFormularioDiceSuFormato:
    def test_el_mensual_dice_su_formato(self, app, bd_copia_real):
        assert codigo_de_formato(PruebaMensualBraq(_UsuarioFalso())) == CODIGO_ACTIVIDAD

    def test_la_redundante_dice_el_mismo_formato_de_actividad(self, app, bd_copia_real):
        assert codigo_de_formato(CalRedundanteFuente(_UsuarioFalso())) == CODIGO_ACTIVIDAD

    def test_la_linealidad_dice_su_propio_formato(self, app, bd_copia_real):
        assert codigo_de_formato(Linealidad(_UsuarioFalso())) == CODIGO_LINEALIDAD

    def test_el_diario_sigue_diciendo_su_formato(self, app, bd_copia_real):
        assert codigo_de_formato(PruebaDiariaBraq(_UsuarioFalso())) == CODIGO_DIARIO


# ---------------------------------------------------------------------------
# (b) Lo que recibe el PDF (generador espiado, sin abrir el visor).
# ---------------------------------------------------------------------------

class TestElPdfRecibeElFormato:
    def _espiar_guardarpdf(self, monkeypatch):
        llamadas = []
        monkeypatch.setattr(
            braq_mensual_mod, "guardarPDF_mensual",
            lambda *args, **kwargs: llamadas.append((args, kwargs)))
        return llamadas

    def test_el_pdf_del_mensual_recibe_el_formato_de_actividad(self, app, bd_copia_real, monkeypatch):
        llamadas = self._espiar_guardarpdf(monkeypatch)
        obj = PruebaMensualBraq(_UsuarioFalso())

        obj.btn_submit.click()

        assert len(llamadas) == 1
        assert llamadas[0][1]["id_maquina"] == CODIGO_ACTIVIDAD

    def test_el_pdf_de_la_redundante_recibe_el_formato_de_actividad(self, app, bd_copia_real, monkeypatch):
        llamadas = self._espiar_guardarpdf(monkeypatch)
        obj = CalRedundanteFuente(_UsuarioFalso())

        obj.btn_submit.click()

        assert len(llamadas) == 1
        assert llamadas[0][1]["id_maquina"] == CODIGO_ACTIVIDAD

    def test_el_pdf_de_linealidad_recibe_el_formato_de_linealidad(self, app, bd_copia_real, monkeypatch):
        recibidos = []

        class _Buffer:
            def getvalue(self):
                return b"%PDF-falso"

        def _espia(*args, **kwargs):
            recibidos.append(kwargs)
            return _Buffer()

        class _Visor:
            def __init__(self, *args, **kwargs):
                pass

            def show(self):
                pass

        class _DbFalsa:
            def close(self):
                pass

        monkeypatch.setattr(braq_mod, "generar_reporte_pdf_multitabla_mensual", _espia)
        monkeypatch.setattr(braq_mod, "PdfViewer", _Visor)
        monkeypatch.setattr(braq_mod, "resource_path", lambda ruta: ruta, raising=False)

        lin = Linealidad(_UsuarioFalso())
        lin.construir_tablas_reporte_linealidad = lambda fecha: {"tabla": object()}
        lin.opeenDatabase = lambda: _DbFalsa()
        lin._obtener_firma_usuario = lambda usuario, db: (None, "Físico Médico")

        lin._generar_reporte_linealidad()

        assert len(recibidos) == 1
        assert recibidos[0]["id_maquina"] == CODIGO_LINEALIDAD


# ---------------------------------------------------------------------------
# (c) Lo viejo sigue igual: el diario y los demás formatos.
# ---------------------------------------------------------------------------

class TestLosDemasFormatosNoCambian:
    def test_censo_exacto_de_codigos_por_hoja_de_encabezado(self):
        assert _codigos_del_excel() == CODIGOS_POR_HOJA

    def test_el_censo_lee_los_dos_widgets_de_un_libro_sintetico(self, tmp_path):
        """El lector del censo cuenta solo la PRIMERA columna `nombres` y
        toma el texto de `descripcion`: libro con dos bloques (como
        `preguntas_mensualBraqui`) y una hoja sin códigos."""
        ruta = str(tmp_path / "sintetico.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "encabezado_x"
        ws.append([None, "nombres", "widget_type", "descripcion", None, "nombres", "widget_type", "descripcion"])
        ws.append([0, "code_1", "QLabel", "AAA", 0, "code_1", "QLabel", "DUPLICADO"])
        ws.append([1, "code_2", "QLabel", "V.9", 1, "code_2", "QLabel", "V.0"])
        wb.create_sheet("sin_codigos").append([None, "nombres", "descripcion"])
        wb.save(ruta)

        assert _codigos_del_excel(ruta) == {"encabezado_x": ("AAA", "V.9")}
