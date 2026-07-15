"""Tests del lector de hojas TRS-398 del OIEA (Fase D3).

Dos capas:

1. Fixture SINTÉTICO (.xlsx generado en el momento con openpyxl): reproduce
   el layout fijo de la plantilla poniendo valores conocidos en las celdas
   mapeadas. Es determinista y corre en cualquier máquina sin depender de
   archivos externos ni de datos clínicos reales.

2. Validación OPCIONAL contra los .xls/.xlsm reales del físico (se salta sola
   con skipif si la carpeta ~/Documents/Archivos_UseApp no está): comprueba el
   descifrado VelvetSweatshop, el mapeo real y que el motor de la app
   reproduce los factores calculados por el Excel oficial.
"""
import os

import openpyxl
import pytest

from services.trs398_excel import (
    leer_trs398, comparar_trs398, CELDAS_ENTRADAS, CELDAS_CALCULADAS,
    CELDAS_ENTRADAS_ELECTRONES, CELDAS_CALCULADAS_ELECTRONES,
    _ref_a_indices, _normalizar, _detectar_tipo_haz,
)

# Caso realista: entradas del Halcyon con sus valores calculados por el Excel.
# Sirve para probar comparar_trs398 sin depender del archivo real.
DATOS_HALCYON = {
    "archivo": "test.xls",
    "entradas": {
        "T_clinica": 21.1, "P_clinica": 85.13, "T0": 22.0, "P0": 101.325,
        "Mplus": 5.525, "Mminus": 5.6, "M1_recomb": 5.525, "M2_recomb": 5.447,
        "a0": 1.198, "a1": -0.8753, "a2": 0.6773,
        "lectura_V1": 5.525, "unidades_monitor": 200.0,
        "tpr2010": 0.626, "factor_calibracion": 0.3034, "PDD_zref": 100.0,
    },
    "calculados": {
        "ktp": 1.1866096830385666, "kpol": 1.006787330316742,
        "ks": 1.0070023695467674, "Mq": 0.03323367808333775,
        "kQ": 0.9964, "Dzref": 0.010046798777934927,
        "dosis_maxima": 0.010046798777934927,
    },
}

CARPETA_REAL = os.path.expanduser("~/Documents/Archivos_UseApp")
HALCYON = os.path.join(CARPETA_REAL, "TRS-398 6 MV FFF Halcyon Dmax.xls")
ELECTRONES_12MEV = os.path.join(
    CARPETA_REAL, "Archivos QA/2024/Enero/iX/Electrones/TRS-398 12 MeV.xls")

# Caso realista de ELECTRONES (E5, auditoría 2026-07-10): hoja real
# Enero/iX/12 MeV, cámara Roos (N34001). Sirve para probar comparar_trs398
# sin depender del corpus real -- la validación contra el archivo real vive
# en TestArchivoRealElectrones12MeV más abajo.
DATOS_ELECTRONES_12MEV = {
    "archivo": "test_electrones.xls",
    "tipo_haz": "electrones",
    "entradas": {
        "r50_medido": 5.127,
        "T_clinica": 21.9, "P_clinica": 85.43, "T0": 20.0, "P0": 101.325,
        "Mplus": 20.96, "Mminus": 20.93, "M1_recomb": 20.96, "M2_recomb": 20.52,
        "a0": 1.022, "a1": -0.3632, "a2": 0.3413,
        "lectura_V1": 21.36, "unidades_monitor": 200.0,
        "factor_calibracion": 0.08563, "PDD_zref": 99.3,
    },
    "calculados": {
        "ktp": 1.1937446812282109, "kpol": 0.9992843511450381,
        "ks": 1.0071056560613143, "Mq": 0.12830595800293096,
        "kQ": 0.9102745360000001, "Dzref": 0.010001039940131951,
        "dosis_maxima": 0.01007154072520841,
        # H3.2: D11/I11 reales de esta misma hoja (Enero/iX/12 MeV) --
        # 1.029*5.127-0.06 y 0.6*Q-0.1, exactos (esta hoja no tiene el
        # override manual de zref que sí aparece en 6 MeV).
        "beam_quality_r50": 5.215682999999999,
        "zref": 3.0294097999999994,
    },
}


# ── Capa 1: fixture sintético determinista ────────────────────────────────

# Valores testigo por clave lógica (números fáciles de rastrear en asserts).
ENTRADAS_TESTIGO = {k: float(i + 1) for i, k in enumerate(CELDAS_ENTRADAS)}
CALCULADOS_TESTIGO = {k: round(1.0 + i / 100, 3) for i, k in enumerate(CELDAS_CALCULADAS)}


@pytest.fixture
def xlsx_sintetico(tmp_path):
    """Construye un .xlsx con la hoja 'Sheet' y los valores testigo en las
    celdas exactas del mapeo, más un '#NAME?' para probar la normalización."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for clave, valor in ENTRADAS_TESTIGO.items():
        ws[CELDAS_ENTRADAS[clave]] = valor
    for clave, valor in CALCULADOS_TESTIGO.items():
        ws[CELDAS_CALCULADAS[clave]] = valor
    # Simular una celda con error de fórmula (como el .xlsm real en kQ)
    ws[CELDAS_CALCULADAS["kQ"]] = "#NAME?"
    ruta = tmp_path / "sintetico.xlsx"
    wb.save(ruta)
    return str(ruta)


class TestMapeoBasico:
    def test_lee_todas_las_entradas(self, xlsx_sintetico):
        datos = leer_trs398(xlsx_sintetico)
        for clave, esperado in ENTRADAS_TESTIGO.items():
            assert datos["entradas"][clave] == esperado, clave

    def test_lee_todos_los_calculados(self, xlsx_sintetico):
        datos = leer_trs398(xlsx_sintetico)
        for clave, esperado in CALCULADOS_TESTIGO.items():
            if clave == "kQ":
                continue  # sobrescrito por #NAME? en el fixture
            assert datos["calculados"][clave] == esperado, clave

    def test_error_de_formula_se_normaliza_a_none(self, xlsx_sintetico):
        datos = leer_trs398(xlsx_sintetico)
        assert datos["calculados"]["kQ"] is None

    def test_incluye_nombre_de_archivo(self, xlsx_sintetico):
        assert leer_trs398(xlsx_sintetico)["archivo"] == "sintetico.xlsx"


class TestManejoDeErrores:
    def test_archivo_inexistente(self):
        with pytest.raises(FileNotFoundError):
            leer_trs398("/no/existe.xls")

    def test_extension_no_soportada(self, tmp_path):
        p = tmp_path / "x.csv"
        p.write_text("a,b")
        with pytest.raises(ValueError, match="Formato no soportado"):
            leer_trs398(str(p))


class TestHelpers:
    @pytest.mark.parametrize("ref, esperado", [
        ("A1", (0, 0)), ("I43", (42, 8)), ("G70", (69, 6)), ("J48", (47, 9)),
    ])
    def test_ref_a_indices(self, ref, esperado):
        assert _ref_a_indices(ref) == esperado

    @pytest.mark.parametrize("entrada, esperado", [
        (None, None), ("", None), ("#NAME?", None), ("#REF!", None),
        ("  water  ", "water"), (1.5, 1.5), (0, 0),
    ])
    def test_normalizar(self, entrada, esperado):
        assert _normalizar(entrada) == esperado


class TestComparacion:
    def test_con_modelo_todo_comparable_y_ok(self):
        filas = comparar_trs398(DATOS_HALCYON, modelo_camara="N31010")
        assert len(filas) == len(CELDAS_CALCULADAS)
        for f in filas:
            assert f["comparable"], f["magnitud"]
            assert f["ok"], f"{f['magnitud']} difiere: {f['diferencia_rel']}"
            assert f["diferencia_rel"] < 0.001

    def test_sin_modelo_kq_y_dependientes_no_comparables(self):
        filas = {f["magnitud"]: f for f in comparar_trs398(DATOS_HALCYON)}
        # Sin modelo de cámara no hay kQ, y sin kQ no hay Dzref ni dosis
        assert filas["kQ"]["comparable"] is False
        assert filas["Dzref"]["comparable"] is False
        assert filas["dosis_maxima"]["comparable"] is False
        # Los que no dependen de kQ sí se comparan
        for clave in ("ktp", "kpol", "ks", "Mq"):
            assert filas[clave]["comparable"] is True
            assert filas[clave]["ok"] is True

    def test_detecta_discrepancia_fuera_de_tolerancia(self):
        datos = {
            "archivo": "x", "entradas": dict(DATOS_HALCYON["entradas"]),
            "calculados": dict(DATOS_HALCYON["calculados"]),
        }
        datos["calculados"]["ktp"] = 1.25  # ~5 % de diferencia, muy fuera
        filas = {f["magnitud"]: f for f in comparar_trs398(datos, modelo_camara="N31010")}
        assert filas["ktp"]["comparable"] is True
        assert filas["ktp"]["ok"] is False
        assert filas["ktp"]["diferencia_rel"] > 0.001

    def test_excel_con_valor_none_no_es_comparable(self):
        datos = {
            "archivo": "x", "entradas": dict(DATOS_HALCYON["entradas"]),
            "calculados": dict(DATOS_HALCYON["calculados"]),
        }
        datos["calculados"]["kQ"] = None  # como el .xlsm real (#NAME?)
        filas = {f["magnitud"]: f for f in comparar_trs398(datos, modelo_camara="N31010")}
        assert filas["kQ"]["comparable"] is False
        assert filas["kQ"]["app"] is not None  # la app sí lo calcula
        assert filas["kQ"]["excel"] is None

    def test_entrada_faltante_deja_magnitud_no_comparable(self):
        datos = {
            "archivo": "x", "entradas": dict(DATOS_HALCYON["entradas"]),
            "calculados": dict(DATOS_HALCYON["calculados"]),
        }
        datos["entradas"]["Mplus"] = None  # sin M+ no hay kpol
        filas = {f["magnitud"]: f for f in comparar_trs398(datos, modelo_camara="N31010")}
        assert filas["kpol"]["app"] is None
        assert filas["kpol"]["comparable"] is False


# ── E5 (2026-07-10): detección de tipo de haz + hojas de ELECTRONES ──────

ENTRADAS_TESTIGO_ELECTRONES = {k: float(i + 1) for i, k in enumerate(CELDAS_ENTRADAS_ELECTRONES)}
CALCULADOS_TESTIGO_ELECTRONES = {
    k: round(1.0 + i / 100, 3) for i, k in enumerate(CELDAS_CALCULADAS_ELECTRONES)}


def _xlsx_con_haz(tmp_path, nombre, texto_a2, entradas_map=None, calculadas_map=None):
    """Construye un .xlsx con la hoja 'Sheet', A2 = texto_a2 (para probar
    _detectar_tipo_haz) y, si se pasan mapas, los valores testigo en sus
    celdas -- mismo patrón que xlsx_sintetico pero parametrizado por haz."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws["A2"] = texto_a2
    if entradas_map:
        for clave, valor in ENTRADAS_TESTIGO_ELECTRONES.items():
            ws[entradas_map[clave]] = valor
    if calculadas_map:
        for clave, valor in CALCULADOS_TESTIGO_ELECTRONES.items():
            ws[calculadas_map[clave]] = valor
    ruta = tmp_path / nombre
    wb.save(ruta)
    return str(ruta)


class TestDeteccionTipoHaz:
    """leer_trs398 debe elegir el mapa de celdas correcto según A2."""

    def test_texto_photon_beam(self, tmp_path):
        ruta = _xlsx_con_haz(tmp_path, "f.xlsx",
                             "in a high-energy photon-beam")
        assert leer_trs398(ruta)["tipo_haz"] == "fotones"

    def test_texto_electron_beam(self, tmp_path):
        ruta = _xlsx_con_haz(tmp_path, "e.xlsx", "in an electron-beam",
                             CELDAS_ENTRADAS_ELECTRONES, CELDAS_CALCULADAS_ELECTRONES)
        datos = leer_trs398(ruta)
        assert datos["tipo_haz"] == "electrones"
        for clave, esperado in ENTRADAS_TESTIGO_ELECTRONES.items():
            assert datos["entradas"][clave] == esperado, clave

    def test_a2_irreconocible_hace_fallback_a_fotones(self, tmp_path):
        """Preserva el comportamiento pre-E5: un archivo sin A2 reconocible
        (como el fixture sintético xlsx_sintetico, que no la setea) se lee
        con el mapa de fotones -- el único que existía antes de esta fase."""
        ruta = _xlsx_con_haz(tmp_path, "vacio.xlsx", "")
        assert leer_trs398(ruta)["tipo_haz"] == "fotones"

    def test_detectar_tipo_haz_directo(self):
        from services.trs398_excel import _LectorCeldas
        lector = _LectorCeldas(lambda f, c: "in an electron-beam" if (f, c) == (1, 0) else None)
        assert _detectar_tipo_haz(lector) == "electrones"


class TestComparacionElectrones:
    """Espejo de TestComparacion pero con la hoja de ELECTRONES (Roos,
    N34001) -- valores centinela de la hoja real Enero/iX 12 MeV."""

    def test_con_modelo_todo_comparable_y_ok(self):
        filas = comparar_trs398(DATOS_ELECTRONES_12MEV, modelo_camara="N34001")
        assert len(filas) == len(CELDAS_CALCULADAS_ELECTRONES)
        for f in filas:
            assert f["comparable"], f["magnitud"]
            assert f["ok"], f"{f['magnitud']} difiere: {f['diferencia_rel']}"

    def test_sin_modelo_kq_y_dependientes_no_comparables(self):
        filas = {f["magnitud"]: f for f in comparar_trs398(DATOS_ELECTRONES_12MEV)}
        assert filas["kQ"]["comparable"] is False
        assert filas["Dzref"]["comparable"] is False
        assert filas["dosis_maxima"]["comparable"] is False
        # beam_quality_r50/zref (H3.2) NO dependen de la cámara -- solo del
        # R50 medido -- así que siguen comparables sin modelo seleccionado.
        for clave in ("ktp", "kpol", "ks", "Mq", "beam_quality_r50", "zref"):
            assert filas[clave]["comparable"] is True
            assert filas[clave]["ok"] is True

    def test_camara_de_fotones_no_tiene_kq_de_electrones(self):
        """N31010 tiene fila en la tabla de FOTONES pero no en la de
        electrones -- comparar una hoja de electrones con esa cámara
        seleccionada debe dejar kQ no comparable, no interpolar la tabla
        equivocada (el bug real que motivó E2/E3)."""
        filas = {f["magnitud"]: f
                for f in comparar_trs398(DATOS_ELECTRONES_12MEV, modelo_camara="N31010")}
        assert filas["kQ"]["comparable"] is False

    def test_incluye_beam_quality_r50_y_zref(self):
        """H3.2 (auditoría 2026-07-14): invierte la exclusión original de E5
        (test_no_incluye_zref_ni_r50w). Aunque 15/73 hojas de electrones del
        corpus 2024 sobreescriben zref a mano (siempre 6 MeV, convención
        clínica -- ver comentario en CELDAS_CALCULADAS_ELECTRONES), la
        decisión pasó a ser mostrar y explicar (aviso en la UI), no ocultar
        -- mismo principio que H3.1."""
        magnitudes = {f["magnitud"] for f in comparar_trs398(DATOS_ELECTRONES_12MEV)}
        assert "beam_quality_r50" in magnitudes
        assert "zref" in magnitudes


# ── Capa 2: validación opcional contra archivos reales ────────────────────

@pytest.fixture(scope="module")
def datos_halcyon():
    return leer_trs398(HALCYON)


@pytest.mark.skipif(not os.path.exists(HALCYON),
                    reason="archivos reales del físico no disponibles en esta máquina")
class TestArchivoRealHalcyon:
    @pytest.fixture
    def datos(self, datos_halcyon):
        return datos_halcyon

    def test_descifra_y_lee_identificacion(self, datos):
        assert datos["entradas"]["acelerador"] == "Halcyon"
        assert datos["entradas"]["serie_camara"] == 1825.0
        assert datos["entradas"]["tpr2010"] == 0.626

    def test_entradas_crudas_esperadas(self, datos):
        e = datos["entradas"]
        assert e["T_clinica"] == 21.1
        assert e["P_clinica"] == 85.13
        assert e["Mplus"] == 5.525
        assert e["Mminus"] == 5.6
        assert e["factor_calibracion"] == 0.3034

    def test_valores_calculados_del_excel(self, datos):
        c = datos["calculados"]
        assert round(c["ktp"], 4) == 1.1866
        assert round(c["kpol"], 5) == 1.00679
        assert round(c["ks"], 4) == 1.007
        assert c["kQ"] == 0.9964
        assert round(c["Dzref"], 6) == 0.010047

    def test_motor_de_la_app_reproduce_al_excel(self, datos):
        """El corazón de D3: recalcular con el motor de la app y confirmar
        que coincide con los factores que el Excel oficial ya trae."""
        from services.dosis_service import DosisService
        e, c = datos["entradas"], datos["calculados"]

        assert DosisService.factor_tp(
            e["T_clinica"], e["P_clinica"], e["T0"], e["P0"]) == round(c["ktp"], 4)
        assert DosisService.factor_k_polaridad(
            e["Mplus"], e["Mminus"]) == round(c["kpol"], 5)
        m1m2 = DosisService.cociente_M1M2(e["M1_recomb"], e["M2_recomb"])
        assert DosisService.Ks_factor(e["a0"], e["a1"], e["a2"], m1m2) == round(c["ks"], 4)
        assert DosisService.interpolar_kq0("N31010", e["tpr2010"]) == c["kQ"]


@pytest.fixture(scope="module")
def datos_electrones_12mev():
    return leer_trs398(ELECTRONES_12MEV)


@pytest.mark.skipif(not os.path.exists(ELECTRONES_12MEV),
                    reason="corpus 2024 no disponible en esta máquina")
class TestArchivoRealElectrones12MeV:
    """Espejo de TestArchivoRealHalcyon para la hoja real de ELECTRONES
    (E5, auditoría 2026-07-10): Enero/iX/Electrones/TRS-398 12 MeV.xls,
    cámara Roos serie 1069 (N34001)."""

    @pytest.fixture
    def datos(self, datos_electrones_12mev):
        return datos_electrones_12mev

    def test_tipo_haz_detectado(self, datos):
        assert datos["tipo_haz"] == "electrones"

    def test_descifra_y_lee_identificacion(self, datos):
        assert datos["entradas"]["acelerador"] == "IX"
        assert datos["entradas"]["serie_camara"] == 1069.0
        assert datos["entradas"]["r50_medido"] == 5.127

    def test_entradas_crudas_esperadas(self, datos):
        e = datos["entradas"]
        assert e["T_clinica"] == 21.9
        assert e["P_clinica"] == 85.43
        assert e["Mplus"] == 20.96
        assert e["Mminus"] == 20.93
        assert e["factor_calibracion"] == 0.08563

    def test_valores_calculados_del_excel(self, datos):
        c = datos["calculados"]
        assert round(c["ktp"], 4) == 1.1937
        assert round(c["kpol"], 5) == 0.99928
        assert round(c["ks"], 4) == 1.0071
        assert round(c["kQ"], 5) == 0.91027
        assert round(c["Dzref"], 6) == 0.010001

    def test_motor_de_la_app_reproduce_al_excel(self, datos):
        """Mismo espíritu que D3 pero para electrones: recalcular con el
        motor de la app y confirmar que coincide con lo que ya trae el
        Excel oficial (E1-E3 corrigieron los 3 bugs que lo impedían)."""
        from services.dosis_service import DosisService
        e, c = datos["entradas"], datos["calculados"]

        assert DosisService.factor_tp(
            e["T_clinica"], e["P_clinica"], e["T0"], e["P0"]) == round(c["ktp"], 4)
        assert DosisService.factor_k_polaridad(
            e["Mplus"], e["Mminus"]) == round(c["kpol"], 5)
        m1m2 = DosisService.cociente_M1M2(e["M1_recomb"], e["M2_recomb"])
        assert DosisService.Ks_factor(e["a0"], e["a1"], e["a2"], m1m2) == round(c["ks"], 4)
        r50w = DosisService.r50_quality(e["r50_medido"])
        assert DosisService.interpolar_r50("N34001", r50w) == round(c["kQ"], 5)


# Las 4 energías de electrones de un mismo mes/máquina (H3.2) -- corpus 2024,
# Junio/iX. OJO: el plan original asumía "los 4 dan Q/zref exactos"; verificado
# contra el archivo real que NO es así -- 6 MeV trae el override clínico de
# zref=1.4 igual que casi todos los demás meses (ver comentario en
# CELDAS_CALCULADAS_ELECTRONES). Esta clase documenta el patrón real, no la
# suposición inicial.
CARPETA_JUNIO_ELECTRONES = os.path.join(
    CARPETA_REAL, "Archivos QA/2024/Junio/IX/Electrones")
ENERGIAS_JUNIO = {
    "6 MeV":  ("TRS-398 6 MeV.xls", False),   # False = zref NO exacto (override 1.4)
    "9 MeV":  ("TRS-398 9 MeV-.xls", True),
    "12 MeV": ("TRS-398 12 MeV.xls", True),
    "15 MeV": ("TRS-398 15 MeV.xls", True),
}


@pytest.mark.skipif(not os.path.isdir(CARPETA_JUNIO_ELECTRONES),
                    reason="corpus 2024 no disponible en esta máquina")
class TestArchivoRealElectronesJunioLas4Energias:
    """Q(R50) exacto en las 4 energías; zref exacto en 9/12/15 MeV pero NO en
    6 MeV (convención clínica fija de esta institución, confirmada en
    prácticamente todos los meses de 2024 -- no es un bug ni una excepción
    aislada de este mes)."""

    @pytest.mark.parametrize("energia,exacto_zref",
                             [(e, ok) for e, (_, ok) in ENERGIAS_JUNIO.items()])
    def test_beam_quality_r50_siempre_exacto(self, energia, exacto_zref):
        archivo, _ = ENERGIAS_JUNIO[energia]
        datos = leer_trs398(os.path.join(CARPETA_JUNIO_ELECTRONES, archivo))
        r50 = datos["entradas"]["r50_medido"]
        q_calc = 1.029 * r50 - 0.06
        assert abs(datos["calculados"]["beam_quality_r50"] - q_calc) < 0.001

    @pytest.mark.parametrize("energia,exacto_zref",
                             [(e, ok) for e, (_, ok) in ENERGIAS_JUNIO.items()])
    def test_zref_exacto_salvo_override_clinico_6mev(self, energia, exacto_zref):
        archivo, _ = ENERGIAS_JUNIO[energia]
        datos = leer_trs398(os.path.join(CARPETA_JUNIO_ELECTRONES, archivo))
        r50 = datos["entradas"]["r50_medido"]
        q_calc = 1.029 * r50 - 0.06
        zref_calc = 0.6 * q_calc - 0.1
        zref_hoja = datos["calculados"]["zref"]
        if exacto_zref:
            assert abs(zref_hoja - zref_calc) < 0.001, (
                f"{energia}: se esperaba fórmula exacta, no un override")
        else:
            assert zref_hoja == 1.4, (
                f"{energia}: se esperaba el override clínico conocido (1.4)")
            assert abs(zref_hoja - zref_calc) > 0.01, (
                f"{energia}: el override debería diferir de la fórmula "
                f"(si ahora coincide, esta hoja ya no es el caso documentado)")
