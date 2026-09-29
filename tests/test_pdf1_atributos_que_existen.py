"""P.1 (PLAN_VOLTAJES_EDITABLES_BRAQUI_28-09.md, segunda ronda, versión
mínima -- tercera ronda §R.3 fila 9): ningún `hasattr(self, 'X')` /
`getattr(self, 'X', ...)` LITERAL puede preguntar por un atributo que
PRODUCCIÓN no crea -- es la clase de defecto que dejaba al PDF del
mensual de braquiterapia pidiéndose siempre como "Control Mensual":
`es_cambio_fuente` (ninguna asignación en todo el árbol) y `cambio_cal`
(el widget real del botón "Cambio de Fuente" es `red_cal`, columna
`nombres` de `widgets.xlsx`) -- los dos solo existían porque un test los
fabricaba a mano (`test_c1_lectura_por_clave_completa.py`), nunca la
pantalla real.

Alcance MÍNIMO, por decisión del físico (§R.3): solo los dos archivos de
braquiterapia (`braq_mensual.py`, `braquiterapia.py`), no todo
`ui/`/`data/`/`models/` -- eso era el alcance del `CAMBIO` original de
`P.1`, y queda fuera de esta ronda.

Censo: para cada llamada literal en los dos archivos, se busca en TODO
el árbol de producción (excluidos `tests/`, `.venv`, `build`,
`__pycache__`) una asignación real `self.X = ...` (incluida una tupla de
asignación múltiple), un método `def X(self...)`, un `setattr(self,
'X', ...)` literal, o una fila con ese nombre en la PRIMERA columna
`nombres` de alguna hoja de `widgets.xlsx`. Si ninguna de las cuatro
fuentes lo explica, el sitio queda en `EXENTAS` con su motivo (censal,
como `ES1`/`RT1`), o el test falla nombrando el archivo, la línea y el
nombre exactos.

**Solo la PRIMERA columna `nombres`**: `preguntas_mensualBraqui` trae un
bloque duplicado a la derecha (columnas T-X) que `createInterface` no
carga (`row['nombres']` es la primera). Ahí vive `cambio_cal` -- medido
sobre las pantallas construidas: `hasattr(pantalla, 'cambio_cal')` es
`False` y `hasattr(pantalla, 'red_cal')` es `True`. Contar ese bloque
habría dado por bueno justo el nombre que motivó esta tarea."""
import ast
import os

import openpyxl
import pytest

ARCHIVOS_ALCANCE = [
    "ui/paginasControles/PruebasMensuales/braq_mensual.py",
    "ui/paginasControles/PruebasDiarias/braquiterapia.py",
]

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DIRECTORIOS_EXCLUIDOS = {".venv", "build", "__pycache__", ".git", "tests"}

WIDGETS_XLSX = os.path.join(RAIZ, "data", "widgets.xlsx")

# Sin excepciones en el alcance mínimo: los 4 sitios que la motivaron
# (es_cambio_fuente x2, cambio_cal x2) se retiraron en el propio CAMBIO
# de P.1. Si algo aparece aquí, o es un defecto real (arreglarlo) o una
# excepción legítima con su motivo escrito (añadirla, nunca en silencio).
EXENTAS = {}


def _sitios_literales(ruta_abs):
    with open(ruta_abs, encoding="utf-8") as f:
        arbol = ast.parse(f.read(), filename=ruta_abs)
    sitios = []
    for nodo in ast.walk(arbol):
        if isinstance(nodo, ast.Call) and isinstance(nodo.func, ast.Name):
            if nodo.func.id in ("hasattr", "getattr") and len(nodo.args) >= 2:
                primero, segundo = nodo.args[0], nodo.args[1]
                if (isinstance(primero, ast.Name) and primero.id == "self"
                        and isinstance(segundo, ast.Constant)
                        and isinstance(segundo.value, str)):
                    sitios.append((nodo.func.id, segundo.value, nodo.lineno))
    return sitios


def _atributos_de_targets(target, destino):
    if isinstance(target, ast.Attribute) and isinstance(target.value, ast.Name) and target.value.id == "self":
        destino.add(target.attr)
    elif isinstance(target, (ast.Tuple, ast.List)):
        for elt in target.elts:
            _atributos_de_targets(elt, destino)
    elif isinstance(target, ast.Starred):
        _atributos_de_targets(target.value, destino)


def _atributos_creados_en(ruta_abs):
    creados = set()
    try:
        with open(ruta_abs, encoding="utf-8", errors="ignore") as f:
            arbol = ast.parse(f.read(), filename=ruta_abs)
    except SyntaxError:
        return creados
    for nodo in ast.walk(arbol):
        if isinstance(nodo, ast.Assign):
            for t in nodo.targets:
                _atributos_de_targets(t, creados)
        elif isinstance(nodo, ast.AnnAssign):
            _atributos_de_targets(nodo.target, creados)
        elif isinstance(nodo, ast.AugAssign):
            _atributos_de_targets(nodo.target, creados)
        elif isinstance(nodo, (ast.FunctionDef, ast.AsyncFunctionDef)):
            creados.add(nodo.name)
        elif isinstance(nodo, ast.Call) and isinstance(nodo.func, ast.Name) and nodo.func.id == "setattr":
            if (len(nodo.args) >= 2 and isinstance(nodo.args[1], ast.Constant)
                    and isinstance(nodo.args[1].value, str)):
                creados.add(nodo.args[1].value)
    return creados


def _nombres_de_widgets_xlsx(ruta):
    """Nombres que `createInterface` crea con `setattr`: la PRIMERA columna
    `nombres` de cada hoja (un bloque duplicado a la derecha no se carga)."""
    nombres = set()
    wb = openpyxl.load_workbook(ruta, data_only=True)
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        encabezado = [c.value for c in ws[1]]
        primera = next((i for i, h in enumerate(encabezado) if h == "nombres"), None)
        if primera is None:
            continue
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if primera < len(fila) and isinstance(fila[primera], str) and fila[primera].strip():
                nombres.add(fila[primera].strip())
    return nombres


def _atributos_creados_en_produccion():
    creados = set()
    for raiz, dirs, files in os.walk(RAIZ):
        dirs[:] = [d for d in dirs if d not in DIRECTORIOS_EXCLUIDOS]
        for nombre in files:
            if nombre.endswith(".py"):
                creados |= _atributos_creados_en(os.path.join(raiz, nombre))
    creados |= _nombres_de_widgets_xlsx(WIDGETS_XLSX)
    return creados


def _censo():
    sitios_por_nombre = {}
    for relativo in ARCHIVOS_ALCANCE:
        for kind, nombre, lineno in _sitios_literales(os.path.join(RAIZ, relativo)):
            sitios_por_nombre.setdefault(nombre, []).append((relativo, lineno, kind))
    return sitios_por_nombre


class TestNingunAtributoPreguntadoEsInexistente:
    def test_todo_sitio_literal_resuelve_o_esta_en_exentas(self):
        censo = _censo()
        creados = _atributos_creados_en_produccion()

        sin_resolver = []
        for nombre, sitios in censo.items():
            if nombre in creados:
                continue
            if nombre in EXENTAS:
                continue
            sin_resolver.append((nombre, sitios))

        assert not sin_resolver, (
            "hasattr/getattr(self, 'X') literal sin ninguna asignación, "
            "método, setattr ni fila de widgets.xlsx que explique 'X' -- "
            "añadir a EXENTAS con motivo si es legítimo, o corregir la "
            f"pantalla si no lo es: {sin_resolver}")

    def test_exentas_no_tiene_entradas_sobrantes(self):
        censo = _censo()
        for nombre in EXENTAS:
            assert nombre in censo, (
                f"'{nombre}' está en EXENTAS pero ya no aparece en ningún "
                "hasattr/getattr literal de los archivos en alcance -- "
                "retirar la excepción, ya no hace falta")

    def test_es_cambio_fuente_y_cambio_cal_ya_no_se_preguntan(self):
        """La causa concreta de P.1 (§1.11): estos dos nombres no los crea
        nada en producción y motivaron esta tarea -- deben haber
        desaparecido de los dos archivos en alcance, no solo quedar en
        EXENTAS."""
        censo = _censo()
        assert "es_cambio_fuente" not in censo
        assert "cambio_cal" not in censo


class TestElCensoDiscriminaDeVerdad:
    def test_un_atributo_inventado_se_detecta(self, monkeypatch):
        import sys
        modulo = sys.modules[__name__]
        censo_original = _censo

        def _censo_con_intruso():
            real = censo_original()
            real.setdefault("atributo_que_nadie_crea_jamas", []).append(
                (ARCHIVOS_ALCANCE[0], 1, "hasattr"))
            return real

        monkeypatch.setattr(modulo, "_censo", _censo_con_intruso)

        censo = modulo._censo()
        creados = modulo._atributos_creados_en_produccion()
        sin_resolver = [n for n in censo if n not in creados and n not in EXENTAS]
        assert "atributo_que_nadie_crea_jamas" in sin_resolver

    def test_solo_cuenta_la_primera_columna_nombres_del_excel(self, tmp_path):
        """El caso de `cambio_cal`: un nombre que solo existe en el bloque
        DUPLICADO de la derecha de una hoja (el que `createInterface` no
        carga) no lo crea nadie. Libro sintético con el mismo trazado que
        `preguntas_mensualBraqui`: dos bloques con encabezado `nombres`."""
        ruta = str(tmp_path / "widgets_sintetico.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "hoja"
        ws.append([None, "nombres", "widget_type", None, "nombres", "widget_type"])
        ws.append([0, "boton_real", "QRadioButton", 0, "boton_duplicado", "QRadioButton"])
        wb.create_sheet("sin_nombres").append(["a", "b"])
        wb.save(ruta)

        nombres = _nombres_de_widgets_xlsx(ruta)

        assert nombres == {"boton_real"}

    def test_el_excel_real_no_cuenta_cambio_cal_pero_si_red_cal(self):
        """Sobre el `widgets.xlsx` de producción: `red_cal` (bloque
        cargado) cuenta; `cambio_cal` (bloque duplicado) no."""
        nombres = _nombres_de_widgets_xlsx(WIDGETS_XLSX)
        assert "red_cal" in nombres
        assert "cambio_cal" not in nombres
